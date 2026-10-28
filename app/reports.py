from __future__ import annotations

from io import BytesIO, StringIO
import json
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from econ_sim import ScenarioComparison, SimulationResult


CORE_SIMULATION_FRAMES = (
    "yearly_summary",
    "monthly_macro",
    "household_income_monthly",
    "government_monthly",
    "company_events",
    "market_events",
    "product_events",
    "products_final",
)


def _safe_scalar(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _mapping_frame(mapping: dict[str, Any], key_name: str = "Setting", value_name: str = "Value") -> pd.DataFrame:
    return pd.DataFrame([{key_name: key, value_name: _safe_scalar(value)} for key, value in mapping.items()])


def _style_workbook(writer: pd.ExcelWriter) -> None:
    wb = writer.book
    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
        # Bounded autofit keeps generated reports readable without huge columns.
        for idx, column_cells in enumerate(ws.columns, start=1):
            width = 10
            for cell in list(column_cells)[:250]:
                if cell.value is not None:
                    width = max(width, min(42, len(str(cell.value)) + 2))
            ws.column_dimensions[get_column_letter(idx)].width = width


def simulation_xlsx_bytes(result: SimulationResult) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        _mapping_frame(result.summary(), "Metric", "Value").to_excel(writer, sheet_name="Summary", index=False)
        _mapping_frame(result.config.to_dict()).to_excel(writer, sheet_name="Config", index=False)
        frames = result.frames()
        for name in CORE_SIMULATION_FRAMES:
            frame = frames.get(name)
            if frame is not None and not frame.empty:
                frame.to_excel(writer, sheet_name=name[:31], index=False)
        _style_workbook(writer)
    return buffer.getvalue()


def simulation_csv_zip_bytes(result: SimulationResult) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("summary.json", json.dumps(result.summary(), indent=2))
        archive.writestr("config.json", json.dumps(result.config.to_dict(), indent=2, default=str))
        frames = result.frames()
        for name in CORE_SIMULATION_FRAMES:
            frame = frames.get(name)
            if frame is not None and not frame.empty:
                archive.writestr(f"{name}.csv", frame.to_csv(index=False))
        archive.writestr(
            "README.txt",
            "Economic Market Simulator report\n"
            "\nThis ZIP contains core CSV outputs from one simulation run.\n"
            "product_monthly is intentionally omitted because long simulations can create very large files.\n",
        )
    return buffer.getvalue()


def _comparison_summary_frame(comparison: ScenarioComparison) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for scenario, values in comparison.summaries().items():
        row = {"Scenario": scenario}
        row.update(values)
        rows.append(row)
    return pd.DataFrame(rows)


def comparison_xlsx_bytes(comparison: ScenarioComparison) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        comparison.kpi_table().to_excel(writer, sheet_name="KPI_Comparison", index=False)
        comparison.kpi_long().to_excel(writer, sheet_name="KPI_Long", index=False)
        _comparison_summary_frame(comparison).to_excel(writer, sheet_name="Scenario_Summaries", index=False)
        _mapping_frame(comparison.baseline.config.to_dict()).to_excel(writer, sheet_name="Baseline_Config", index=False)
        for name, result in comparison.scenarios.items():
            safe = "".join(ch for ch in name if ch.isalnum() or ch in " _-").strip() or "Scenario"
            _mapping_frame(result.config.to_dict()).to_excel(writer, sheet_name=(f"Cfg_{safe}")[:31], index=False)
        for metric, rows in comparison.yearly_dashboard().items():
            pd.DataFrame(rows).to_excel(writer, sheet_name=(f"Yearly_{metric}")[:31], index=False)
        _style_workbook(writer)
    return buffer.getvalue()


def comparison_csv_zip_bytes(comparison: ScenarioComparison) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("kpi_comparison.csv", comparison.kpi_table().to_csv(index=False))
        archive.writestr("kpi_long.csv", comparison.kpi_long().to_csv(index=False))
        archive.writestr("scenario_summaries.csv", _comparison_summary_frame(comparison).to_csv(index=False))
        archive.writestr("baseline_config.json", json.dumps(comparison.baseline.config.to_dict(), indent=2, default=str))
        for name, result in comparison.scenarios.items():
            safe = "".join(ch for ch in name if ch.isalnum() or ch in " _-").strip().replace(" ", "_") or "scenario"
            archive.writestr(f"{safe}_config.json", json.dumps(result.config.to_dict(), indent=2, default=str))
        for metric, rows in comparison.yearly_dashboard().items():
            archive.writestr(f"yearly_{metric}.csv", pd.DataFrame(rows).to_csv(index=False))
        archive.writestr(
            "README.txt",
            "Economic Market Simulator scenario-comparison report\n"
            "\nThe comparison uses the same baseline random seed unless a scenario explicitly overrides it.\n",
        )
    return buffer.getvalue()
